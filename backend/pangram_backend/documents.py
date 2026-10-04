from io import BytesIO
from pathlib import Path
import zipfile
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text
from PIL import Image, UnidentifiedImageError
from .network import request_bytes


class DocumentError(ValueError):
    pass


def extract_document(filename: str, data: bytes, max_chars=500_000):
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv"}:
            text = data.decode("utf-8-sig")
        elif suffix == ".rtf":
            if not data.startswith(b"{\\rtf"):
                raise DocumentError("Not a valid RTF document")
            text = rtf_to_text(data.decode("latin-1"))
        elif suffix == ".docx":
            with zipfile.ZipFile(BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > 10000 or sum(m.file_size for m in members) > 50_000_000:
                    raise DocumentError("Document expands beyond the safe size limit")
            doc = Document(BytesIO(data))
            paragraphs = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                paragraphs.extend("\t".join(c.text for c in row.cells) for row in table.rows)
            text = "\n\n".join(paragraphs)
        elif suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise DocumentError("Not a valid PDF document")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise DocumentError("Password-protected PDFs are not supported")
            if len(reader.pages) > 300:
                raise DocumentError("PDFs are limited to 300 pages")
            parts, length = [], 0
            for page in reader.pages:
                part = page.extract_text() or ""
                length += len(part)
                if length > max_chars:
                    raise DocumentError("Extracted text exceeds the character limit")
                parts.append(part)
            text = "\n\n".join(parts)
            if not text.strip():
                raise DocumentError("This PDF has no extractable text. OCR is not configured; upload a text-based document")
        else:
            raise DocumentError("Supported documents: TXT, MD, CSV, RTF, DOCX, PDF")
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError("The file is damaged, incorrectly labeled, or cannot be read") from e
    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentError("The document has no text")
    if len(text) > max_chars:
        raise DocumentError("Extracted text exceeds the character limit")
    return text


def fetch_article(url):
    data, content_type, final_url = request_bytes(url, max_bytes=2_000_000)
    if not any(t in content_type.lower() for t in ["text/html", "text/plain", "application/xhtml+xml"]):
        raise DocumentError("URL did not return a text or HTML document")
    if "text/plain" in content_type:
        return data.decode("utf-8", errors="replace"), final_url
    soup = BeautifulSoup(data, "html.parser")
    for element in soup(["script", "style", "nav", "header", "footer", "aside", "noscript"]):
        element.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    return root.get_text("\n", strip=True), final_url


def validate_image(data):
    try:
        with Image.open(BytesIO(data)) as image:
            if image.format not in {"JPEG", "PNG", "WEBP"}:
                raise DocumentError("Supported images: JPEG, PNG, WebP")
            if image.width < 512 or image.height < 512:
                raise DocumentError("Image dimensions must be at least 512 × 512")
            if image.width * image.height > 40_000_000:
                raise DocumentError("Image exceeds the 40 megapixel limit")
            result = {"width": image.width, "height": image.height, "format": image.format}
            image.verify()
            return result
    except DocumentError:
        raise
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as e:
        raise DocumentError("The image could not be read safely") from e
